"""Reads the workshop's tracking workbook into plain blocks (one per product), untouched.
Usage: python3 -I parse.py source.xlsx out_dir  -> out_dir/raw.json + out_dir/img/<sheet>-<row>.png"""
import sys, json, datetime, os, re
import openpyxl

SIZES = ['XS', 'S', 'M', 'L', 'XL']
path, out = sys.argv[1], sys.argv[2]
os.makedirs(os.path.join(out, 'img'), exist_ok=True)
wb = openpyxl.load_workbook(path, data_only=True)
raw = []
for ws in wb.worksheets:
    if ws.title == 'SHEET MẪU':
        continue
    head = {c.column: str(c.value).strip() for c in ws[2] if c.value}
    top = {c.column: str(c.value).strip() for c in ws[1] if c.value}
    # Delivery size columns sit between "SỐ LƯỢNG GIAO KHÁCH" and "NGÀY GIAO".
    start = next(k for k, v in top.items() if v == 'SỐ LƯỢNG GIAO KHÁCH')
    date_col = next(k for k, v in top.items() if v == 'NGÀY GIAO')
    note_col = next(k for k, v in top.items() if v == 'GHI CHÚ')
    ship_cols = {col: head[col] for col in range(start, date_col) if col in head}
    images = {}
    for n, im in enumerate(ws._images):
        row = im.anchor._from.row + 1
        name = f'{ws.title}-{row}.png'
        with open(os.path.join(out, 'img', name), 'wb') as f:
            f.write(im._data())
        images[row] = name
    last = max((r[0].row for r in ws.iter_rows(min_row=3) if any(c.value is not None for c in r[:date_col])), default=2)
    def val(r, c):
        return ws.cell(row=r, column=c).value
    starts = [r for r in range(3, last + 1) if str(val(r, 5) or '').strip() in ('XS', 'NO SIZE') or (val(r, 2) and str(val(r, 5) or '').strip() not in SIZES + ['TỔNG'])]
    starts = sorted(set(starts))
    for i, r0 in enumerate(starts):
        r1 = (starts[i + 1] - 1) if i + 1 < len(starts) else last
        block = {'sheet': ws.title, 'row': r0, 'end': r1, 'name': (str(val(r0, 2)).strip() if val(r0, 2) else ''), 'date': None, 'sizes': [], 'ships': [], 'notes': [], 'images': [images[r] for r in sorted(images) if r0 <= r <= r1]}
        c = val(r0, 3)
        block['date'] = c.strftime('%Y-%m-%d') if isinstance(c, datetime.datetime) else (str(c).strip() if c else None)
        for r in range(r0, r1 + 1):
            size = str(val(r, 5) or '').strip()
            if size and size != 'TỔNG':
                block['sizes'].append({'size': size, 'ordered': val(r, 4), 'cut': val(r, 6)})
            d = val(r, date_col)
            qty = {ship_cols[col]: int(val(r, col)) for col in ship_cols if isinstance(val(r, col), (int, float)) and val(r, col)}
            if qty:
                block['ships'].append({'date': d.strftime('%Y-%m-%d') if isinstance(d, datetime.datetime) else (str(d).strip() if d else None), 'qty': qty})
            note = val(r, note_col)
            if note:
                block['notes'].append(str(note).strip())
        raw.append(block)
json.dump(raw, open(os.path.join(out, 'raw.json'), 'w'), ensure_ascii=False, indent=1)
for b in raw:
    o = sum(int(s['ordered'] or 0) for s in b['sizes'])
    d = sum(sum(s['qty'].values()) for s in b['ships'])
    print(f"{b['sheet']:9} r{b['row']:<4} {b['name'][:38]!r:42} date={b['date']} ordered={o} delivered={d} ships={len(b['ships'])} img={b['images']} notes={b['notes']}")
print(len(raw), 'blocks')
